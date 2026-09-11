from __future__ import annotations

import io
import re
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook

HEADERS = [
    "Company ID",
    "Company Name",
    "Legal Name",
    "Country",
    "City",
    "Address",
    "Website",
    "Telephone",
    "Nature of Business",
    "Company Status",
    "Notes",
    "Owner Email",
]

STATUSES = {"Prospect", "Customer", "Former Customer", "Inactive"}


def _norm_name(value: str) -> str:
    folded = value.strip().casefold()
    folded = re.sub(r"[^\w\s]", "", folded, flags=re.UNICODE)
    return re.sub(r"\s+", " ", folded).strip()


def _cell(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _as_uuid(value: str | None) -> str | None:
    if not value:
        return None
    try:
        return str(UUID(value))
    except ValueError:
        return None


def workbook_bytes(rows: list[dict[str, Any]]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Companies"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row.get(key) or "" for key in HEADERS])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def read_rows(payload: bytes) -> list[dict[str, str | None]]:
    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    sheet = book.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        raise ValueError("Excel file has no header row")
    names = [str(cell).strip() if cell is not None else "" for cell in header]
    missing = [name for name in HEADERS if name not in names]
    if missing:
        raise ValueError(f"Excel is missing columns: {', '.join(missing)}")
    index = {name: names.index(name) for name in HEADERS}
    parsed: list[dict[str, str | None]] = []
    for raw in rows:
        if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
            continue
        parsed.append({key: _cell(raw[index[key]] if index[key] < len(raw) else None) for key in HEADERS})
    return parsed


def classify_row(
    raw: dict[str, str | None],
    by_id: dict[str, dict[str, Any]],
    by_name: dict[tuple[str, str], dict[str, Any]],
    members: dict[str, str],
) -> dict[str, Any]:
    name = raw.get("Company Name")
    country = (raw.get("Country") or "").upper() or None
    status = raw.get("Company Status")
    owner_email = raw.get("Owner Email")
    company_id = _as_uuid(raw.get("Company ID"))
    error = None
    if status and status not in STATUSES:
        error = "invalid Company Status"
    elif country and len(country) != 2:
        error = "Country must be a 2-letter code"
    owner_id = members.get(owner_email.casefold()) if owner_email else None
    if owner_email and not owner_id:
        error = "Owner Email is not a member of this entity"
    match = by_id.get(company_id or "")
    if not match and name:
        match = by_name.get((_norm_name(name), country or ""))
        if not match and country:
            match = by_name.get((_norm_name(name), ""))
    action = "skip"
    if error:
        action = "skip"
    elif match:
        action = "update"
        company_id = str(match["id"])
    elif name:
        action = "create"
    else:
        error = "Company Name is required"
        action = "skip"
    return {
        "action": action,
        "error": error,
        "company_id": company_id,
        "company_name": name,
        "legal_name": raw.get("Legal Name"),
        "country": country,
        "city": raw.get("City"),
        "address": raw.get("Address"),
        "website": raw.get("Website"),
        "telephone": raw.get("Telephone"),
        "nature_of_business": raw.get("Nature of Business"),
        "status": status if status in STATUSES else None,
        "notes": raw.get("Notes"),
        "owner_user_id": owner_id,
    }
